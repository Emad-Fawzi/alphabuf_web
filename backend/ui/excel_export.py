"""
تصدير نتائج الفحص لملف Excel منسّق (تلوين القرارات، تنسيق الأرقام، تجميد الصف الأول،
فلتر تلقائي). مستقل عن أي واجهة عرض - بيرجع bytes جاهزة للتحميل.
"""
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from utils.formatting import fmt_thousands


def apply_excel_formatting(df_res):
    """بيرجع bytes لملف Excel منسّق من DataFrame نتائج الفحص."""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_res.to_excel(writer, index=False, sheet_name="Screener_Results")
        wb = writer.book
        ws = writer.sheets["Screener_Results"]
        ws.sheet_view.rightToLeft = True  # عشان الأعمدة والنصوص العربية المختلطة بالإنجليزي تترتب صح جوه إكسل

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        thin_border = Border(*(Side(style="thin", color="D1D5DB"),) * 4)

        col_index = {name: i + 1 for i, name in enumerate(df_res.columns)}
        explanation_col_nums = {col_index.get("شرح السوينغ"), col_index.get("شرح نصيحة المالك"),
                                 col_index.get("استراتيجية الدخول والخروج")} - {None}

        buy_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        sell_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        wait_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

        for row_num, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row,
                                                     min_col=1, max_col=ws.max_column), start=1):
            for cell in row:
                col_num = cell.column
                cell.border = thin_border
                if row_num == 1:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    continue

                if col_num in explanation_col_nums:
                    cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                if row_num > 1:
                    if col_num == col_index.get("القيمة"):
                        cell.number_format = '#,##0'
                    elif col_num in (col_index.get("السعر"), col_index.get("الدعم"), col_index.get("المقاومة"),
                                      col_index.get("سعر شراء مقترح"), col_index.get("سعر بيع مقترح"),
                                      col_index.get("وقف الخسارة"), col_index.get("هدف جني الأرباح")):
                        cell.number_format = '0.00'
                    elif col_num == col_index.get("التغير %"):
                        cell.number_format = '+0.00"%";-0.00"%"'
                    elif col_num in (col_index.get("قرار السوينغ"), col_index.get("نصيحة المالك")):
                        text = str(cell.value)
                        if any(k in text for k in ["اشترِ", "استمر", "احتفظ"]):
                            cell.fill = buy_fill
                        elif any(k in text for k in ["لا تشترِ", "اخرج", "جني أرباح"]):
                            cell.fill = sell_fill
                        elif "انتظر" in text or "اصبر" in text:
                            cell.fill = wait_fill

        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
            col_letter = get_column_letter(col_cells[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 45)

        for row_num in range(2, ws.max_row + 1):
            ws.row_dimensions[row_num].height = 60

        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
    return output.getvalue()
